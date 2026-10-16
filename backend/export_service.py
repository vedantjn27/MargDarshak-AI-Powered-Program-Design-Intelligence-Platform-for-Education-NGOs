"""Readable, organization-specific program documents."""

from html import escape
from io import BytesIO
from pathlib import Path
from tempfile import NamedTemporaryFile

import matplotlib.pyplot as plt
import networkx as nx
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.util import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, KeepTogether, PageBreak, Table, TableStyle
from reportlab.platypus import Image as PDFImage
from PIL import Image as PillowImage, ImageDraw, ImageFont
from presentation import presentable_text
from export_visuals import available_visuals, visual_notes, theory_of_change as render_toc_visual, _font as visual_font
from template_catalog import canonical_theme


def _path(suffix):
    with NamedTemporaryFile(prefix="margdarshak_", suffix=suffix, delete=False) as file:
        return file.name


HINDI_LABELS = {
    "organization_profile": "संगठन परिचय", "problem_definition": "समस्या परिभाषा",
    "problem_tree": "समस्या वृक्ष", "outcomes": "छात्र परिणाम",
    "methodology": "कार्यपद्धति", "theory_of_change": "परिवर्तन का सिद्धांत",
    "stakeholders": "हितधारक", "practice_changes": "अभ्यास परिवर्तन",
    "indicators": "संकेतक", "indicator_baselines": "आधार और लक्ष्य",
    "theme": "विषय", "geography": "स्थान", "state": "राज्य", "district": "ज़िला",
    "core_problem": "मुख्य समस्या", "affected_group": "प्रभावित समूह",
    "root_causes": "मूल कारण", "effects": "प्रभाव", "smart_outcomes": "मापने योग्य परिणाम",
    "selected": "चुनी गई पद्धतियाँ", "custom_components": "अतिरिक्त घटक",
    "activities": "गतिविधियाँ", "outputs": "तत्काल परिणाम", "impact": "दीर्घकालिक प्रभाव",
    "baseline_value": "आधार मूल्य", "current_value": "वर्तमान मूल्य",
    "target_value": "लक्ष्य मूल्य", "indicator_name": "संकेतक नाम",
    "start_date": "आरंभ तिथि", "end_date": "समाप्ति तिथि",
}


def _label(value, language="en"):
    return HINDI_LABELS.get(str(value), str(value).replace("_", " ").title()) if language == "hi" else str(value).replace("_", " ").title()


def _lines(value, prefix="", language="en"):
    """Turn nested snapshot content into readable labels and statements."""
    if isinstance(value, dict):
        for key, item in value.items():
            if key.endswith("_id") or key in {"id", "nodes", "edges", "mermaid_diagram", "organization_id", "template_id"}:
                continue
            heading = _label(key, language)
            if isinstance(item, (dict, list)):
                yield from _lines(item, f"{prefix}{heading}: ", language)
            elif item is not None and presentable_text(item):
                yield f"{prefix}{heading}: {presentable_text(item)}"
    elif isinstance(value, list):
        for item in value:
            if isinstance(item, (dict, list)):
                yield from _lines(item, prefix, language)
            else:
                yield f"{prefix}{presentable_text(item)}"
    elif value is not None:
        yield f"{prefix}{presentable_text(value)}"


def _sections(snapshot, language="en"):
    order = ["organization_profile", "problem_definition", "problem_tree", "outcomes", "methodology", "theory_of_change", "stakeholders", "practice_changes", "indicators", "indicator_baselines"]
    for key in order:
        value = snapshot.get(key)
        if value:
            lines = [line for line in _lines(value, language=language) if line.strip()]
            if lines:
                yield _label(key, language), lines


def _hindi_block(text, size=12, color="#1e293b"):
    """Render shaped Devanagari text for ReportLab's PDF canvas."""
    font = ImageFont.truetype(str(Path(__file__).parent / "assets" / "NotoSansDevanagari.ttf"), size * 2, layout_engine=ImageFont.Layout.RAQM)
    measure = PillowImage.new("RGB", (980, 10), "white")
    drawer = ImageDraw.Draw(measure)
    lines, current = [], ""
    for word in str(text).split():
        candidate = f"{current} {word}".strip()
        if drawer.textlength(candidate, font=font) > 950 and current:
            lines.append(current)
            current = word
        else:
            current = candidate
    if current:
        lines.append(current)
    lines = lines or [" "]
    line_height = size * 3
    picture = PillowImage.new("RGB", (980, line_height * len(lines) + 12), "white")
    draw = ImageDraw.Draw(picture)
    for index, line in enumerate(lines):
        draw.text((0, index * line_height), line, font=font, fill=color)
    buffer = BytesIO()
    picture.save(buffer, format="PNG")
    buffer.seek(0)
    return PDFImage(buffer, width=490, height=picture.height / 2)


def pdf(snapshot, organization, language="en"):
    path = _path(".pdf")
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="AtlasTitle", parent=styles["Title"], textColor=colors.HexColor("#0b3041"), fontSize=31, leading=38, spaceAfter=14))
    styles.add(ParagraphStyle(name="AtlasSection", parent=styles["Heading2"], textColor=colors.HexColor("#0b3041"), fontSize=19, spaceAfter=12))
    styles.add(ParagraphStyle(name="AtlasBody", parent=styles["BodyText"], textColor=colors.HexColor("#294357"), fontSize=10, leading=15, spaceAfter=8))
    visuals = available_visuals(snapshot, language)
    name = organization.get("organization_name", "Organization")
    profile = snapshot.get("organization_profile") or {}
    focus = canonical_theme(profile.get("theme") or (organization.get("thematic_focus") or [""])[0]).replace("_", " ").title()
    geography = organization.get("geography") or {}
    location = ", ".join(filter(None, [geography.get("district"), geography.get("state")]))
    story = []

    def add_text(value, style="AtlasBody", *, size=11, color="#294357"):
        if language == "hi":
            story.append(_hindi_block(value, size, color))
        else:
            story.append(Paragraph(escape(str(value)), styles[style]))

    add_text("कार्यक्रम डिज़ाइन एटलस" if language == "hi" else "PROGRAM DESIGN ATLAS", "AtlasTitle", size=25, color="#0b6f78")
    add_text(name, "AtlasSection", size=19, color="#0b3041")
    if location:
        add_text(location, size=11)
    add_text(focus, size=11, color="#0b6f78")
    story.append(Spacer(1, 16))
    art = _cover_path(profile.get("theme") or (organization.get("thematic_focus") or [""])[0])
    if art:
        story.append(PDFImage(str(art), width=490, height=276))
        story.append(Spacer(1, 18))

    problem = (snapshot.get("problem_definition") or {}).get("core_problem") or (snapshot.get("problem_tree") or {}).get("core_problem")
    outcome = (snapshot.get("outcomes") or {}).get("outcome_statement")
    method = ", ".join((snapshot.get("methodology") or {}).get("selected") or [])
    for title, value in (
        (("The need", "समस्या")[language == "hi"], problem),
        (("The intended change", "अपेक्षित बदलाव")[language == "hi"], outcome),
        (("The approach", "कार्यपद्धति")[language == "hi"], method),
    ):
        if value:
            add_text(title, "AtlasSection", size=14, color="#0b6f78")
            add_text(presentable_text(value), size=11)

    if visuals:
        story.append(PageBreak())
    for title, figure in visuals:
        image = PillowImage.open(BytesIO(figure))
        scale = min(490 / image.width, 550 / image.height)
        heading = _hindi_block(title, 16, "#0b3041") if language == "hi" else Paragraph(escape(title), styles["AtlasSection"])
        caption_text = "Based on the saved design; confirm the underlying evidence before sharing." if language == "en" else "सहेजी गई योजना पर आधारित; साझा करने से पहले प्रमाण जाँचें।"
        caption = _hindi_block(caption_text, 9, "#64748b") if language == "hi" else Paragraph(escape(caption_text), styles["AtlasBody"])
        notes = []
        for item in visual_notes(snapshot, title, language):
            words = presentable_text(item).split()
            if words:
                notes.append(" ".join(words[:22]) + ("…" if len(words) > 22 else ""))
        note_blocks = [
            _hindi_block("• " + note, 10, "#294357") if language == "hi" else Paragraph("<font color='#0b6f78'>•</font> " + escape(note), styles["AtlasBody"])
            for note in notes[:2 if language == "hi" else 3]
        ]
        story.append(KeepTogether([
            heading,
            PDFImage(BytesIO(figure), width=image.width * scale, height=image.height * scale),
            Spacer(1, 8), *note_blocks, caption, Spacer(1, 18),
        ]))

    notes = list(_sections(snapshot, language))
    if notes:
        story.append(PageBreak())
        add_text("Design notes" if language == "en" else "योजना के विवरण", "AtlasSection", size=18, color="#0b3041")
        for title, lines in notes:
            if not lines:
                continue
            heading = _hindi_block(title, 14, "#0b6f78") if language == "hi" else Paragraph(escape(title), styles["AtlasSection"])
            first = _hindi_block(lines[0], 10) if language == "hi" else Paragraph(escape(lines[0]), styles["AtlasBody"])
            story.append(KeepTogether([heading, first]))
            for line in lines[1:]:
                add_text(line, size=10)

    def page_frame(canvas, document):
        canvas.saveState()
        canvas.setStrokeColor(colors.HexColor("#17b7a4"))
        canvas.setLineWidth(3)
        canvas.line(48, 813, 547, 813)
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor("#64748b"))
        canvas.drawString(48, 29, "MargDarshak  |  Program Design")
        canvas.drawRightString(547, 29, str(document.page))
        canvas.restoreState()

    SimpleDocTemplate(path, pagesize=(595, 842), rightMargin=48, leftMargin=48, topMargin=48, bottomMargin=48).build(
        story, onFirstPage=page_frame, onLaterPages=page_frame)
    return path


def _cover_path(theme):
    key = canonical_theme(theme)
    paths = (
        Path(__file__).resolve().parent / "assets" / "covers" / f"{key}.png",
        Path(__file__).resolve().parent.parent / "frontend" / "public" / "covers" / f"{key}.png",
    )
    return next((path for path in paths if path.is_file()), None)


def _rgb(value):
    value = value.lstrip("#")
    return RGBColor(*(int(value[index:index + 2], 16) for index in (0, 2, 4)))


def _ppt_rect(slide, x, y, width, height, color, *, radius=False):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(color)
    shape.line.fill.background()
    return shape


def _ppt_fitted_text(value, width, height, size, language, bold=False):
    """Size slide text before writing it so different viewers keep it inside its box."""
    content = presentable_text(value)
    if not content:
        return "", size
    draw = ImageDraw.Draw(PillowImage.new("RGB", (1, 1)))
    usable_width = max(40, int((width - 0.1) * 96 * 0.9))
    usable_height = max(16, int(height * 96 * 0.9))

    def fits(text, font_size):
        font_pixels = max(10, round(font_size * 96 / 72))
        font = visual_font(font_pixels, language, bold)
        rows = 0
        for paragraph in text.split("\n"):
            line = ""
            for word in paragraph.split():
                candidate = f"{line} {word}".strip()
                if line and draw.textlength(candidate, font=font) > usable_width:
                    rows += 1
                    line = word
                else:
                    line = candidate
            rows += max(1, int(draw.textlength(line, font=font) / usable_width + 0.999)) if line else 1
        return rows * round(font_pixels * 1.25) <= usable_height

    smallest_size = min(size, 11)
    for font_size in range(size, smallest_size - 1, -1):
        if fits(content, font_size):
            return content, font_size
    words = content.split()
    low, high = 0, len(words)
    while low < high:
        middle = (low + high + 1) // 2
        shortened = " ".join(words[:middle]) + ("…" if middle < len(words) else "")
        if fits(shortened, smallest_size):
            low = middle
        else:
            high = middle - 1
    return (" ".join(words[:low]) + ("…" if low < len(words) else "")), smallest_size


def _ppt_text(slide, value, x, y, width, height, *, size=20, color="#f7fbfc", bold=False, language="en", align=None):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    frame.margin_left = Inches(0.04)
    frame.margin_right = Inches(0.04)
    frame.margin_top = Inches(0)
    frame.margin_bottom = Inches(0)
    paragraph = frame.paragraphs[0]
    fitted_text, fitted_size = _ppt_fitted_text(value, width, height, size, language, bold)
    paragraph.text = fitted_text
    paragraph.font.name = "Nirmala UI" if language == "hi" else "Aptos"
    paragraph.font.size = Pt(fitted_size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = _rgb(color)
    if align is not None:
        paragraph.alignment = align
    return shape


def _ppt_base(deck, heading, language="en", accent="#17b7a4"):
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = _rgb("#081a29")
    _ppt_rect(slide, 0, 0, 0.13, 7.5, accent)
    _ppt_text(slide, heading, 0.65, 0.31, 11.9, 0.6, size=27, bold=True, language=language)
    _ppt_rect(slide, 0.65, 1.02, 1.8, 0.05, accent)
    _ppt_text(slide, "MARGDARSHak  /  PROGRAM DESIGN", 0.65, 7.1, 8, 0.2, size=9, color="#8eacba")
    _ppt_text(slide, str(len(deck.slides)).zfill(2), 12, 7.06, 0.6, 0.24, size=10, color="#8eacba", align=PP_ALIGN.RIGHT)
    return slide


def _ppt_card(slide, heading, body, x, y, width, height, accent, language="en"):
    _ppt_rect(slide, x, y, width, height, "#102b3c", radius=True)
    _ppt_rect(slide, x, y, 0.08, height, accent)
    _ppt_text(slide, heading, x + 0.22, y + 0.18, width - 0.42, 0.55, size=18, color=accent, bold=True, language=language)
    _ppt_text(slide, body or "—", x + 0.22, y + 0.79, width - 0.42, height - 0.95, size=17, language=language)


def presentation(snapshot, organization, language="en"):
    path = _path(".pptx")
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    profile = snapshot.get("organization_profile") or {}
    theme = profile.get("theme") or (organization.get("thematic_focus") or [""])[0]
    name = organization.get("organization_name", "Organization")
    visuals = available_visuals(snapshot, language)

    cover = _ppt_base(deck, "", language)
    art = _cover_path(theme)
    if art:
        cover.shapes.add_picture(str(art), Inches(5.4), Inches(0), height=Inches(7.5))
    _ppt_rect(cover, 0, 0, 5.9, 7.5, "#081a29")
    _ppt_rect(cover, 0.7, 1.25, 1.7, 0.08, "#17b7a4")
    _ppt_text(cover, "PROGRAM DESIGN" if language == "en" else "कार्यक्रम डिज़ाइन", 0.7, 1.55, 4.6, 0.55, size=21, color="#17b7a4", bold=True, language=language)
    _ppt_text(cover, name, 0.7, 2.16, 4.8, 2.35, size=35, bold=True, language=language)
    _ppt_text(cover, canonical_theme(theme).replace("_", " ").title(), 0.7, 5.24, 4.8, 0.58, size=19, color="#f4ba5b")
    _ppt_text(cover, "MARGDARSHak", 0.7, 6.7, 3.5, 0.28, size=11, color="#bed0d9")

    glance = _ppt_base(deck, "Program at a glance" if language == "en" else "कार्यक्रम एक नज़र में", language)
    problem = (snapshot.get("problem_definition") or {}).get("core_problem") or (snapshot.get("problem_tree") or {}).get("core_problem")
    outcome = (snapshot.get("outcomes") or {}).get("outcome_statement")
    methods = ", ".join((snapshot.get("methodology") or {}).get("selected") or [])
    _ppt_card(glance, "THE NEED" if language == "en" else "समस्या", problem, 0.65, 1.42, 3.9, 4.68, "#f17d68", language)
    _ppt_card(glance, "THE CHANGE" if language == "en" else "बदलाव", outcome, 4.72, 1.42, 3.9, 4.68, "#17b7a4", language)
    _ppt_card(glance, "THE APPROACH" if language == "en" else "तरीका", methods, 8.79, 1.42, 3.9, 4.68, "#f4ba5b", language)

    for title, figure in visuals:
        slide = _ppt_base(deck, title, language)
        image = PillowImage.open(BytesIO(figure))
        scale = min(12.0 / image.width, 5.65 / image.height)
        width, height = image.width * scale, image.height * scale
        slide.shapes.add_picture(BytesIO(figure), Inches((13.333 - width) / 2), Inches(1.28 + (5.65 - height) / 2), width=Inches(width), height=Inches(height))

    stakeholders = snapshot.get("stakeholders") or []
    if stakeholders or methods:
        slide = _ppt_base(deck, "People and delivery" if language == "en" else "लोग और कार्यान्वयन", language, "#f4ba5b")
        _ppt_card(slide, "METHODS" if language == "en" else "कार्यपद्धतियाँ", "\n".join((snapshot.get("methodology") or {}).get("selected") or []), 0.65, 1.45, 5.9, 4.9, "#f4ba5b", language)
        _ppt_card(slide, "STAKEHOLDERS" if language == "en" else "हितधारक", "\n".join(presentable_text(item) for item in stakeholders), 6.8, 1.45, 5.9, 4.9, "#ad9af7", language)

    finish = _ppt_base(deck, "From plan to action" if language == "en" else "योजना से कार्रवाई तक", language, "#ad9af7")
    steps = (
        ("01  VERIFY", "Confirm baseline and barriers with local evidence."),
        ("02  DELIVER", "Agree on owners, resources, and a practical schedule."),
        ("03  LEARN", "Review indicators and adapt when the data changes."),
    ) if language == "en" else (
        ("01  जाँचें", "स्थानीय प्रमाण से आधार रेखा और बाधाएँ जाँचें।"),
        ("02  लागू करें", "जिम्मेदारियाँ, संसाधन और समय तय करें।"),
        ("03  सीखें", "संकेतक देखें और प्रमाण के अनुसार योजना बदलें।"),
    )
    for index, (heading, body) in enumerate(steps):
        _ppt_card(finish, heading, body, 0.65 + index * 4.08, 1.8, 3.8, 3.9, ("#17b7a4", "#f4ba5b", "#ad9af7")[index], language)
    deck.save(path)
    return path


def _sheet_style(sheet):
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="4338CA")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True)
    for column in sheet.columns:
        letter = column[0].column_letter
        sheet.column_dimensions[letter].width = min(65, max(18, max(len(str(cell.value or "")) for cell in column) + 3))
    sheet.auto_filter.ref = sheet.dimensions


def indicators(snapshot, organization, language="en"):
    path = _path(".xlsx")
    book = Workbook()
    sheet = book.active
    sheet.title = "Indicator Plan"
    sheet.append(["संगठन", "संकेतक", "आधार", "वर्तमान", "लक्ष्य", "प्रगति %", "आरंभ", "समाप्ति"] if language == "hi" else ["Organization", "Indicator", "Baseline", "Current", "Target", "Progress %", "Start", "End"])
    targets = snapshot.get("indicator_baselines", [])
    for row in targets:
        baseline = row.get("baseline_value")
        current = row.get("current_value")
        target = row.get("target_value")
        progress = round(max(0, min(100, (current - baseline) / (target - baseline) * 100)), 1) if current is not None and baseline is not None and target is not None and target > baseline else 0
        sheet.append([organization["organization_name"], row.get("indicator_name"), baseline, current, target, progress, row.get("start_date"), row.get("end_date")])
    _sheet_style(sheet)
    book.save(path)
    return path


def budget(snapshot, organization, language="en"):
    path = _path(".xlsx")
    book = Workbook()
    sheet = book.active
    sheet.title = "Budget Planning"
    sheet.append(["संगठन", "पद्धति / घटक", "इकाई लागत", "मात्रा", "अनुमानित कुल", "टिप्पणी"] if language == "hi" else ["Organization", "Methodology / Component", "Unit cost", "Quantity", "Estimated total", "Notes"])
    method = snapshot.get("methodology", {})
    entries = method.get("selected", []) + method.get("custom_components", [])
    for index, item in enumerate(entries, start=2):
        sheet.append([organization["organization_name"], item, None, None, f"=C{index}*D{index}", "स्थानीय लागत और मात्रा भरें" if language == "hi" else "Enter local cost and quantity"])
    _sheet_style(sheet)
    book.save(path)
    return path


def toc_image(snapshot, organization, language="en"):
    path = _path(".png")
    rendered = render_toc_visual(snapshot, language)
    if rendered:
        Path(path).write_bytes(rendered)
        return path
    toc = snapshot.get("theory_of_change", {})
    graph = nx.DiGraph()
    positions = {}
    if toc.get("nodes"):
        levels = {"activity": 0, "output": 1, "outcome": 2, "impact": 3}
        counts = {key: 0 for key in levels}
        for node in toc["nodes"]:
            kind = node.get("type", "activity")
            x = levels.get(kind, 0)
            y = counts.get(kind, 0)
            counts[kind] = y + 1
            graph.add_node(node["id"], label=presentable_text(node.get("label", ""))[:45])
            positions[node["id"]] = (x * 4, -y * 2)
        for edge in toc.get("edges", []):
            if edge.get("source") in graph and edge.get("target") in graph:
                graph.add_edge(edge["source"], edge["target"])
    else:
        layers = [(key, toc.get(key, [])) for key in ("activities", "outputs", "outcomes", "impact")]
        for x, (_, labels) in enumerate(layers):
            for y, label in enumerate(labels):
                node = f"{x}-{y}"
                graph.add_node(node, label=presentable_text(label)[:45])
                positions[node] = (x * 4, -y * 2)
    fig, axis = plt.subplots(figsize=(16, 9))
    axis.set_title(f"{'परिवर्तन का सिद्धांत' if language == 'hi' else 'Theory of Change'} | {organization['organization_name']}", fontsize=18, pad=22)
    if graph:
        nx.draw_networkx(graph, positions, labels=nx.get_node_attributes(graph, "label"), ax=axis,
                         node_color="#c7d2fe", edge_color="#6366f1", node_size=4200, font_size=8, arrows=True)
    else:
        axis.text(0.5, 0.5, "Add activities, outputs, outcomes and impact to build this diagram.", ha="center")
    axis.axis("off")
    fig.tight_layout()
    fig.savefig(path, dpi=180, bbox_inches="tight")
    plt.close(fig)
    return path


def create(export_type, snapshot, organization, language="en"):
    handlers = {"LFA_PDF": pdf, "TOC_IMAGE": toc_image, "INDICATOR_EXCEL": indicators,
                "BUDGET_EXCEL": budget, "PRESENTATION_PPT": presentation}
    return handlers[export_type](snapshot, organization, language)
